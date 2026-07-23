from pathlib import Path

from docx import Document


GUIDE = Path(__file__).resolve().parents[1] / "Student_Test_Package" / "Virtual_Makerspace_Student_Guide_EN.docx"


def replace_paragraph(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


document = Document(GUIDE)

paragraph_updates = {
    4: "Current Build | Version 1.0.7 | July 23, 2026 | Oklahoma-Alabama Two-Person Pilot Guide",
    16: "For the Person Installing the App",
    17: "1) Extract the complete student ZIP. 2) Connect the Quest by USB. 3) Approve USB debugging inside the headset. 4) Run INSTALL_TO_QUEST.cmd. The installer removes older copies, verifies version 1.0.7, and launches the app directly in immersive VR.",
    33: "Screen 4. Check Room, Relay, and Voice Status",
    48: "Both headsets show the same ROOM code, 2/2 CONNECTED, VOICE | MIC ON, and ACTIVITY UNLOCKED.",
    49: "Use CREATE ROOM on the first headset and JOIN ROOM on the second headset. Begin only after both devices confirm the same room code and the activity unlock message.",
    59: "4. Confirm 2/2 CONNECTED, VOICE | MIC ON, and ACTIVITY UNLOCKED | USE TRIGGER TO GRAB PARTS.",
    61: "Live Quest Verification Requirement",
    62: "The automated two-process test has verified one Host and one Guest in the same Session, Relay/Netcode connection, and Vivox channel. After the physical test, run COLLECT_TEST_RESULTS.cmd on both headset computers and send both reports to the instructor.",
}
for index, text in paragraph_updates.items():
    replace_paragraph(document.paragraphs[index], text)

replace_paragraph(
    document.tables[0].cell(0, 0).paragraphs[0],
    "Release Verification  Version 1.0.7 passed 29/29 Unity regression tests and independent Host/Guest Session, Relay/Netcode, and Vivox tests. It also waits for Quest microphone permission before voice startup.",
)
replace_paragraph(
    document.tables[1].cell(0, 0).paragraphs[0],
    "Learner Lobby  English CREATE ROOM and JOIN ROOM controls, room-code entry, live participant status, and voice status are included. At 2/2 Connected, the lobby stops intercepting controller rays and displays ACTIVITY UNLOCKED.",
)

controller_table = document.tables[4]
controller_table.cell(1, 1).text = (
    "Point at the resistor or LED and press the Trigger button. Move the component near a glowing breadboard socket and release Trigger."
)
controller_table.cell(2, 1).text = (
    "The component follows your controller and can be placed at the designated socket."
)

replace_paragraph(
    document.tables[7].cell(0, 0).paragraphs[0],
    "Connection Requirement  Do not begin until both headsets display the same room code, 2/2 CONNECTED, VOICE | MIC ON, and ACTIVITY UNLOCKED.",
)

status_table = document.tables[8]
status_table.cell(0, 1).text = (
    "The world-space lobby shows a ROOM code, connection state, participant count, voice state, and beginner guidance."
)
status_table.cell(1, 1).text = (
    "Learner A selects CREATE ROOM. Learner B enters that exact code and selects JOIN ROOM. Allow microphone access on both Quest headsets."
)
status_table.cell(2, 1).text = (
    "Both headsets show the same code, 2/2 CONNECTED, VOICE | MIC ON, and ACTIVITY UNLOCKED | USE TRIGGER TO GRAB PARTS."
)
status_table.cell(3, 1).text = (
    "Check Wi-Fi and the room code. Rerun the clean installer if the app opens as a 2D panel. Send VirtualMakerspace_Device_Diagnostics.txt if startup still fails."
)

troubleshooting = document.tables[11]
def get_or_add_troubleshooting_row(label):
    for existing_row in troubleshooting.rows:
        if existing_row.cells[0].text.strip() == label:
            return existing_row.cells
    row_cells = troubleshooting.add_row().cells
    row_cells[0].text = label
    return row_cells


row = get_or_add_troubleshooting_row("The app opens as a floating 2D panel")
row[0].text = "The app opens as a floating 2D panel"
row[1].text = "Discard older packages and run the version 1.0.7 INSTALL_TO_QUEST.cmd. It performs a clean install and uses the Quest VR launch category."
row = get_or_add_troubleshooting_row("Player is already a member of the lobby")
row[0].text = "Player is already a member of the lobby"
row[1].text = "Version 1.0.7 cleans the stale membership and retries once. If it repeats, rerun the clean installer and send the diagnostics file."

document.save(GUIDE)
print(GUIDE)
